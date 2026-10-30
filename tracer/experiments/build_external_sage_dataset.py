"""Convert UF Retail, THOR-MAGNI, and HuRoN/SACSoN to SAGE windows.

The converter keeps recording boundaries intact.  Each output split contains
fixed-size windows with absolute scene positions in history and future
velocities as targets.  UF Retail has no robot mocap in the local release, so
its adapter uses an explicitly marked static virtual robot at the initial
human position; it is a forecasting portability experiment, not an
action-conditioned robot experiment.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Iterable

import numpy as np

from tracer.data.heterogeneous_scene import AgentRole

try:
    import openpyxl
except ImportError:  # pragma: no cover - dependency is checked by the CLI.
    openpyxl = None  # type: ignore[assignment]

try:
    from rosbags.rosbag1 import Reader
    from rosbags.typesys import Stores, get_typestore
except ImportError:  # pragma: no cover - dependency is checked by the CLI.
    Reader = None  # type: ignore[assignment,misc]
    Stores = None  # type: ignore[assignment,misc]
    get_typestore = None  # type: ignore[assignment,misc]


ROLE_ROBOT = int(AgentRole.ROBOT)
ROLE_PHA = int(AgentRole.PHA)
ROLE_HNAA = int(AgentRole.HNAA)
ROLE_SSAA = int(AgentRole.SSAA)
ROLE_UNKNOWN = int(AgentRole.UNKNOWN_ENTITY)


def _number(value: Any) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _resample_tracks(times: np.ndarray, positions: np.ndarray, dt: float) -> tuple[np.ndarray, np.ndarray]:
    """Resample tracks to a regular clock without extrapolating outside support."""

    if len(times) < 2:
        return np.empty(0, dtype=np.float64), np.empty((0, positions.shape[1], 2), dtype=np.float32)
    order = np.argsort(times)
    times = np.asarray(times, dtype=np.float64)[order]
    positions = np.asarray(positions, dtype=np.float64)[order]
    unique, unique_indices = np.unique(times, return_index=True)
    times = unique
    positions = positions[unique_indices]
    grid = np.arange(times[0], times[-1] + dt * 0.25, dt, dtype=np.float64)
    output = np.full((len(grid), positions.shape[1], 2), np.nan, dtype=np.float32)
    for track in range(positions.shape[1]):
        valid = np.isfinite(positions[:, track, :]).all(axis=1)
        if np.sum(valid) < 2:
            continue
        source_t = times[valid]
        for axis in range(2):
            output[:, track, axis] = np.interp(grid, source_t, positions[valid, track, axis]).astype(np.float32)
        output[grid < source_t[0] - 1e-9, track] = np.nan
        output[grid > source_t[-1] + 1e-9, track] = np.nan
    return grid, output


def _read_uf(path: Path, dt: float) -> tuple[str, np.ndarray, np.ndarray] | None:
    if openpyxl is None:
        raise RuntimeError("UF Retail conversion requires openpyxl")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "position" not in workbook.sheetnames:
            return None
        rows: list[tuple[float, float, float]] = []
        sheet = workbook["position"]
        for row in sheet.iter_rows(min_col=1, max_col=3, values_only=True):
            timestamp = _number(row[0])
            x = _number(row[1])
            y = _number(row[2])
            if timestamp is not None and x is not None and y is not None:
                rows.append((timestamp / 1000.0, x, y))
    finally:
        workbook.close()
    if len(rows) < 32:
        return None
    raw = np.asarray(rows, dtype=np.float64)
    times, human = _resample_tracks(raw[:, 0], raw[:, None, 1:3], dt)
    if len(times) < 32:
        return None
    human = human[:, 0, :]
    human -= human[0]
    robot = np.zeros_like(human)
    positions = np.stack([robot, human], axis=1)
    roles = np.asarray([ROLE_ROBOT, ROLE_PHA], dtype=np.int64)
    return path.stem, positions, roles


def _thor_role(role: str) -> int:
    lowered = role.lower()
    if any(token in lowered for token in ("visitor", "carrier", "human", "participant")):
        return ROLE_PHA
    if any(token in lowered for token in ("vehicle", "robot", "cart")):
        return ROLE_HNAA
    if any(token in lowered for token in ("object", "storage", "static")):
        return ROLE_SSAA
    return ROLE_UNKNOWN


def _read_thor(path: Path, dt: float) -> tuple[str, np.ndarray, np.ndarray] | None:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        try:
            metadata_rows = [next(reader) for _ in range(16)]
            header = next(reader)
        except StopIteration:
            return None
        metadata = {row[0]: row[1:] for row in metadata_rows if row and row[0]}
        names = [name for name in metadata.get("BODY_NAMES", []) if name]
        role_values = metadata.get("BODY_ROLES", [])
        roles_by_name = {name: role_values[index] if index < len(role_values) else "" for index, name in enumerate(names)}
        columns: dict[str, tuple[int, int]] = {}
        for name in names:
            xy = tuple(header.index(f"{name} Centroid_{axis}") for axis in ("X", "Y") if f"{name} Centroid_{axis}" in header)
            if len(xy) == 2:
                columns[name] = xy  # type: ignore[assignment]
        robot_name = "DARKO_Robot"
        if robot_name not in columns:
            return None
        ordered_names = [robot_name] + [name for name in names if name != robot_name and name in columns]
        rows: list[list[float]] = []
        times: list[float] = []
        for row in reader:
            if len(row) < len(header):
                row.extend([""] * (len(header) - len(row)))
            timestamp = _number(row[header.index("Time")] if "Time" in header else None)
            if timestamp is None:
                continue
            point = np.full((len(ordered_names), 2), np.nan, dtype=np.float64)
            for index, name in enumerate(ordered_names):
                x = _number(row[columns[name][0]])
                y = _number(row[columns[name][1]])
                if x is not None and y is not None:
                    point[index] = (x / 1000.0, y / 1000.0)
            times.append(timestamp)
            rows.append(point.reshape(-1).tolist())
    if len(rows) < 32:
        return None
    raw_positions = np.asarray(rows, dtype=np.float64).reshape(len(rows), len(ordered_names), 2)
    times_array, positions = _resample_tracks(np.asarray(times), raw_positions, dt)
    if len(times_array) < 32:
        return None
    roles = np.asarray([ROLE_ROBOT] + [_thor_role(roles_by_name.get(name, "")) for name in ordered_names[1:]], dtype=np.int64)
    return path.stem, positions, roles


def _yaw_from_quaternion(quaternion: Any) -> float:
    return float(2.0 * math.atan2(float(quaternion.z), float(quaternion.w)))


def _interp_pose(times: np.ndarray, xy: np.ndarray, yaw: np.ndarray, query: float) -> tuple[float, float, float] | None:
    if len(times) < 2 or query < times[0] or query > times[-1]:
        return None
    return (
        float(np.interp(query, times, xy[:, 0])),
        float(np.interp(query, times, xy[:, 1])),
        float(np.interp(query, times, yaw)),
    )


def _read_huron(path: Path, dt: float) -> tuple[str, np.ndarray, np.ndarray] | None:
    if Reader is None or Stores is None or get_typestore is None:
        raise RuntimeError("HuRoN conversion requires rosbags")
    typestore = get_typestore(Stores.ROS1_NOETIC)
    odom_times: list[float] = []
    odom_xy: list[tuple[float, float]] = []
    odom_yaw: list[float] = []
    marker_rows: dict[int, list[tuple[float, float, float, str]]] = {}
    try:
        with Reader(path) as reader:
            connections = {connection.topic: connection for connection in reader.connections}
            selected = [connection for topic in ("/odometry", "/pedestrians_pose") if (connection := connections.get(topic))]
            if not selected:
                return None
            for connection, timestamp_ns, raw in reader.messages(connections=selected):
                timestamp = float(timestamp_ns) / 1e9
                message = typestore.deserialize_ros1(raw, connection.msgtype)
                if connection.topic == "/odometry":
                    position = message.pose.pose.position
                    x = _number(position.x)
                    y = _number(position.y)
                    if x is None or y is None:
                        continue
                    odom_times.append(timestamp)
                    odom_xy.append((x, y))
                    odom_yaw.append(_yaw_from_quaternion(message.pose.pose.orientation))
                else:
                    for marker in message.markers:
                        if int(marker.action) == 2:
                            continue
                        x = _number(marker.pose.position.x)
                        y = _number(marker.pose.position.y)
                        if x is None or y is None:
                            continue
                        marker_rows.setdefault(int(marker.id), []).append(
                            (timestamp, x, y, str(marker.header.frame_id))
                        )
    except Exception:
        return None
    if len(odom_times) < 32:
        return None
    order = np.argsort(np.asarray(odom_times))
    odom_t = np.asarray(odom_times, dtype=np.float64)[order]
    odom_xy_array = np.asarray(odom_xy, dtype=np.float64)[order]
    odom_yaw_array = np.unwrap(np.asarray(odom_yaw, dtype=np.float64)[order])
    grid, robot_track = _resample_tracks(odom_t, odom_xy_array[:, None, :], dt)
    if len(grid) < 32:
        return None
    tracks: list[np.ndarray] = [robot_track[:, 0, :]]
    marker_ids: list[int] = []
    for marker_id, observations in sorted(marker_rows.items()):
        observations.sort(key=lambda item: item[0])
        world_rows: list[tuple[float, float, float]] = []
        for timestamp, rel_x, rel_y, frame_id in observations:
            pose = _interp_pose(odom_t, odom_xy_array, odom_yaw_array, timestamp)
            if pose is None:
                continue
            robot_x, robot_y, yaw = pose
            if frame_id in {"base_link", "base_footprint", ""}:
                x = robot_x + math.cos(yaw) * rel_x - math.sin(yaw) * rel_y
                y = robot_y + math.sin(yaw) * rel_x + math.cos(yaw) * rel_y
            else:
                x, y = rel_x, rel_y
            world_rows.append((timestamp, x, y))
        if len(world_rows) < 8:
            continue
        values = np.asarray(world_rows, dtype=np.float64)
        track = np.full((len(grid), 2), np.nan, dtype=np.float32)
        unique_t, unique_idx = np.unique(values[:, 0], return_index=True)
        values = values[unique_idx]
        for axis in range(2):
            track[:, axis] = np.interp(grid, values[:, 0], values[:, axis + 1]).astype(np.float32)
        inside = (grid >= values[0, 0]) & (grid <= values[-1, 0])
        gaps = np.diff(values[:, 0])
        indices = np.searchsorted(values[:, 0], grid, side="right")
        previous = np.clip(indices - 1, 0, len(values) - 1)
        next_index = np.clip(indices, 0, len(values) - 1)
        valid_gap = np.ones(len(grid), dtype=bool)
        middle = (indices > 0) & (indices < len(values))
        valid_gap[middle] = gaps[indices[middle] - 1] <= 0.6
        track[~(inside & valid_gap)] = np.nan
        tracks.append(track)
        marker_ids.append(marker_id)
    if len(tracks) < 2:
        return None
    positions = np.stack(tracks, axis=1)
    roles = np.asarray([ROLE_ROBOT] + [ROLE_PHA] * len(marker_ids), dtype=np.int64)
    return path.stem, positions, roles


def _split_paths(paths: list[Path]) -> dict[str, list[Path]]:
    paths = sorted(paths)
    if len(paths) < 3:
        return {"train": paths, "val": paths, "test": paths}
    test = [path for index, path in enumerate(paths) if index % 5 == 0]
    val = [path for index, path in enumerate(paths) if index % 5 == 1]
    train = [path for path in paths if path not in set(test) and path not in set(val)]
    return {"train": train, "val": val, "test": test}


def _make_windows(
    sequence_id: str,
    positions: np.ndarray,
    roles: np.ndarray,
    dt: float,
    history: int,
    horizon: int,
    max_entities: int,
    stride: int,
    max_speed: float,
) -> dict[str, np.ndarray]:
    agents = 1 + max_entities
    histories: list[np.ndarray] = []
    futures: list[np.ndarray] = []
    currents: list[np.ndarray] = []
    goals: list[np.ndarray] = []
    role_rows: list[np.ndarray] = []
    agent_masks: list[np.ndarray] = []
    robot_masks: list[np.ndarray] = []
    entity_masks: list[np.ndarray] = []
    loss_masks: list[np.ndarray] = []
    for start in range(0, len(positions) - history - horizon + 1, max(1, stride)):
        end = start + history + horizon
        window = positions[start:end]
        if not np.isfinite(window[:, 0, :]).all():
            continue
        available: list[tuple[float, int]] = []
        current_robot = window[history - 1, 0]
        for index in range(1, positions.shape[1]):
            if np.isfinite(window[:, index, :]).all():
                distance = float(np.linalg.norm(window[history - 1, index] - current_robot))
                available.append((distance, index))
        if not available:
            continue
        available.sort(key=lambda item: (item[0], item[1]))
        selected = [0] + [index for _, index in available[:max_entities]]
        history_positions = window[:history, selected].transpose(1, 0, 2).astype(np.float32)
        future_positions = window[history:, selected].transpose(1, 0, 2).astype(np.float32)
        previous = history_positions[:, -1:, :]
        velocity = np.diff(np.concatenate([previous, future_positions], axis=1), axis=1) / float(dt)
        if max_speed > 0 and float(np.linalg.norm(velocity, axis=-1).max()) > max_speed:
            continue
        history_padded = np.zeros((agents, history, 2), dtype=np.float32)
        future_padded = np.zeros((agents, horizon, 2), dtype=np.float32)
        current_padded = np.zeros((agents, 2), dtype=np.float32)
        goal_padded = np.zeros((agents, 2), dtype=np.float32)
        roles_padded = np.full((agents,), ROLE_UNKNOWN, dtype=np.int64)
        active = np.zeros((agents,), dtype=bool)
        robot = np.zeros((agents,), dtype=bool)
        entity = np.zeros((agents,), dtype=bool)
        loss = np.zeros((agents,), dtype=bool)
        count = len(selected)
        history_padded[:count] = history_positions
        future_padded[:count] = velocity
        current_padded[:count] = history_positions[:, -1]
        goal_padded[:count] = future_positions[:, -1]
        roles_padded[:count] = roles[selected]
        active[:count] = True
        robot[0] = True
        entity[1:count] = True
        loss[:count] = True
        histories.append(history_padded)
        futures.append(future_padded)
        currents.append(current_padded)
        goals.append(goal_padded)
        role_rows.append(roles_padded)
        agent_masks.append(active)
        robot_masks.append(robot)
        entity_masks.append(entity)
        loss_masks.append(loss)
    if not histories:
        raise ValueError(f"no complete windows generated for {sequence_id}")
    return {
        "history": np.stack(histories),
        "future": np.stack(futures),
        "current_pos": np.stack(currents),
        "goals": np.stack(goals),
        "goal_mask": np.asarray(robot_masks, dtype=bool),
        "roles": np.stack(role_rows),
        "agent_mask": np.stack(agent_masks),
        "robot_mask": np.stack(robot_masks),
        "entity_mask": np.stack(entity_masks),
        "loss_mask": np.stack(loss_masks),
        "dt": np.asarray([dt], dtype=np.float32),
    }


def _concat(records: list[dict[str, np.ndarray]], dt: float) -> dict[str, np.ndarray]:
    if not records:
        raise ValueError("no records produced any complete windows")
    result: dict[str, np.ndarray] = {}
    for key in records[0]:
        if key == "dt":
            result[key] = np.asarray([dt], dtype=np.float32)
        else:
            result[key] = np.concatenate([record[key] for record in records], axis=0)
    return result


def _recording_paths(dataset: str, root: Path) -> list[Path]:
    if dataset == "uf_retail":
        return sorted(
            path
            for path in root.glob("*.xlsx")
            if (
                len(path.stem.split("_")) == 2
                and path.stem.split("_")[1].isdigit()
                and len(path.stem.split("_")[1]) == 3
            )
            or path.stem.endswith("_offaxis")
        )
    if dataset == "thor_magni":
        return sorted(root.rglob("*.csv"))
    if dataset == "huron_sacson":
        return sorted(root.rglob("*.bag"))
    raise ValueError(dataset)


def _read_sequence(dataset: str, path: Path, dt: float):
    if dataset == "uf_retail":
        return _read_uf(path, dt)
    if dataset == "thor_magni":
        return _read_thor(path, dt)
    return _read_huron(path, dt)


def convert_dataset(
    dataset: str,
    root: Path,
    output_dir: Path,
    history: int,
    horizon: int,
    max_entities: int,
    dt: float,
    stride: int,
    max_speed: float,
    limit_recordings: int = 0,
) -> dict[str, Any]:
    paths = _recording_paths(dataset, root)
    if limit_recordings > 0:
        paths = paths[:limit_recordings]
    split_paths = _split_paths(paths)
    output_dir.mkdir(parents=True, exist_ok=True)
    summary: dict[str, Any] = {"dataset": dataset, "root": str(root), "dt": dt, "history": history, "horizon": horizon, "max_entities": max_entities, "stride": stride, "max_speed": max_speed, "robot_policy": "virtual_static" if dataset == "uf_retail" else "recorded_or_odometry_transformed", "splits": {}}
    for split, files in split_paths.items():
        records: list[dict[str, np.ndarray]] = []
        used: list[str] = []
        skipped: list[dict[str, str]] = []
        for path in files:
            try:
                parsed = _read_sequence(dataset, path, dt)
                if parsed is None:
                    skipped.append({"path": str(path), "reason": "insufficient_or_invalid_track"})
                    continue
                sequence_id, positions, roles = parsed
                records.append(_make_windows(sequence_id, positions, roles, dt, history, horizon, max_entities, stride, max_speed))
                used.append(sequence_id)
            except Exception as error:  # Keep one corrupt public file from hiding valid recordings.
                skipped.append({"path": str(path), "reason": f"{type(error).__name__}: {error}"})
        arrays = _concat(records, dt)
        output_path = output_dir / f"{dataset}_{split}.npz"
        np.savez_compressed(output_path, **arrays)
        manifest = {
            "dataset": dataset,
            "split": split,
            "source_root": str(root),
            "recordings_seen": len(files),
            "recordings_used": len(used),
            "recording_ids": used,
            "skipped": skipped,
            "windows": int(arrays["history"].shape[0]),
            "history": history,
            "horizon": horizon,
            "max_entities": max_entities,
            "dt": dt,
            "window_stride": stride,
            "max_speed_qc_mps": max_speed,
            "coordinate_policy": "meters; UF centered on first human position; THOR centroid mm converted to meters; HuRoN base_link markers transformed with odometry yaw",
            "robot_policy": "static virtual robot because the local UF mocap release has no robot trajectory" if dataset == "uf_retail" else "recorded robot trajectory",
            "claim_boundary": "forecasting/portability evaluation only; no real-human causal or persistent-effect claim",
        }
        output_path.with_suffix(".json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        summary["splits"][split] = manifest
    digest = hashlib.sha256(json.dumps(summary, sort_keys=True).encode()).hexdigest()
    summary["conversion_id"] = digest[:16]
    (output_dir / f"{dataset}_conversion_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Build unified SAGE NPZ windows for external datasets.")
    parser.add_argument("--dataset", required=True, choices=["uf_retail", "thor_magni", "huron_sacson"])
    parser.add_argument("--root", required=True)
    parser.add_argument("--output-dir", default="experiments/external_sage_data")
    parser.add_argument("--history", type=int, default=8)
    parser.add_argument("--horizon", type=int, default=12)
    parser.add_argument("--max-entities", type=int, default=8)
    parser.add_argument("--dt", type=float, default=0.1)
    parser.add_argument("--stride", type=int, default=4)
    parser.add_argument("--max-speed", type=float, default=5.0)
    parser.add_argument("--limit-recordings", type=int, default=0)
    args = parser.parse_args()
    summary = convert_dataset(
        args.dataset,
        Path(args.root),
        Path(args.output_dir),
        args.history,
        args.horizon,
        args.max_entities,
        args.dt,
        args.stride,
        args.max_speed,
        args.limit_recordings,
    )
    for split, record in summary["splits"].items():
        print(f"{args.dataset} {split}: recordings={record['recordings_used']} windows={record['windows']}")
    print(f"conversion_id={summary['conversion_id']}")


if __name__ == "__main__":
    main()
